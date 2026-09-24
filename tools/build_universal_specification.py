from __future__ import annotations

import json
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]

sys.path.insert(0, str(Path(__file__).resolve().parent))
from spec_counts import (  # noqa: E402
    literal_collection_size,
    spec_docx_name,
    test_method_count,
)

sys.path.pop(0)

OUTPUT = ROOT / spec_docx_name(ROOT)


def _package_version() -> str:
    manifest = json.loads(
        (ROOT / "SOURCE-MANIFEST.json").read_text(encoding="utf-8")
    )
    version = manifest["package_version"]
    if not isinstance(version, str) or not re.fullmatch(r"\d+\.\d+\.\d+", version):
        raise RuntimeError(f"Invalid package_version: {version!r}")
    return version


SPEC_VERSION = _package_version()
SPEC_VERSION_MAJOR_MINOR = ".".join(SPEC_VERSION.split(".")[:2])
SPEC_DATE_TEXT = "23 September 2026"

CONFORMANCE_CHECK_COUNT = literal_collection_size(ROOT / "tests" / "test_contract_conformance.py", "EXPECTED_CONFORMANCE_IDS")
# The two sets are counted apart, because the acceptance report names them
# apart. While the bridge reached the runner through an import inside the
# regressions there was one figure and one gate row; now there are two of each,
# and a single summed sentence would describe a suite that no longer exists.
EXTERNAL_REGRESSION_COUNT = test_method_count(
    ROOT / "tests" / "test_contract_regressions.py", "ContractRegressions"
)
OTCS_BRIDGE_TEST_COUNT = test_method_count(
    ROOT / "tests" / "test_otcs_bridge.py", "OTCSBridgeRegressions"
)
FUNCTIONAL_TEST_COUNT = test_method_count(ROOT / "tests" / "test_universal_connection.py", "ConnectionSuite")
SEMANTIC_TEST_COUNT = test_method_count(ROOT / "tests" / "test_semantic_contracts.py", "SemanticContractTests")
SAFETY_MUTATION_COUNT = literal_collection_size(ROOT / "tests" / "test_safety_mutations.py", "EXPECTED_MUTATION_IDS")


def _normalize_docx_archive(path: Path) -> None:
    """Make package metadata and ZIP member timestamps reproducible."""
    fixed_time = (2026, 7, 28, 0, 0, 0)
    temp = path.with_name(path.stem + ".deterministic" + path.suffix)
    with zipfile.ZipFile(path, "r") as source:
        members = [(info, source.read(info.filename)) for info in source.infolist()]

    rewritten = []
    app_metadata_seen = False
    for info, data in members:
        if info.filename == "docProps/app.xml":
            xml = data.decode("utf-8")
            xml, application_count = re.subn(
                r"<Application>.*?</Application>",
                "<Application>NC2.5 Deterministic Builder</Application>",
                xml,
                count=1,
            )
            xml, version_count = re.subn(
                r"<AppVersion>.*?</AppVersion>",
                "<AppVersion>1.0</AppVersion>",
                xml,
                count=1,
            )
            if application_count != 1 or version_count != 1:
                raise RuntimeError("DOCX application metadata shape drifted")
            for tag in ("Pages", "Words", "Characters", "CharactersWithSpaces", "Lines", "Paragraphs"):
                xml = re.sub(fr"<{tag}>.*?</{tag}>", "", xml, count=1)
            data = xml.encode("utf-8")
            app_metadata_seen = True
        rewritten.append((info, data))

    if not app_metadata_seen:
        raise RuntimeError("DOCX application metadata is missing")

    try:
        with zipfile.ZipFile(temp, "w") as target:
            for info, data in sorted(rewritten, key=lambda item: item[0].filename):
                member = zipfile.ZipInfo(info.filename, fixed_time)
                member.compress_type = info.compress_type
                member.external_attr = info.external_attr
                member.internal_attr = info.internal_attr
                member.create_system = info.create_system
                member.comment = info.comment
                target.writestr(member, data)
        temp.replace(path)
    finally:
        if temp.exists():
            temp.unlink()

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "17365D"
LIGHT_BLUE = "DCE6F1"
PALE_BLUE = "EEF4FA"
LIGHT_GREY = "F2F4F7"
MID_GREY = "D9E1E8"
DARK_GREY = "4A5568"
WHITE = "FFFFFF"
AMBER = "FFF2CC"
RED_LIGHT = "FCE8E6"
GREEN_LIGHT = "E2F0D9"
TOTAL_TABLE_DXA = 9360


def load_json(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: Sequence[int], *, indent_dxa: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TOTAL_TABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for grid_col in list(grid):
        grid.remove(grid_col)
    for width in widths:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)


def set_cell_margins(cell, top: int = 80, start: int = 120, bottom: int = 80, end: int = 120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    table_header = OxmlElement("w:tblHeader")
    table_header.set(qn("w:val"), "true")
    tr_pr.append(table_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def keep_row_with_next(row) -> None:
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.keep_with_next = True


def set_cell_text(
    cell,
    text: str,
    *,
    bold: bool = False,
    color: str = "000000",
    size: float = 8.5,
) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def add_table(
    doc: Document,
    headers: Sequence[str],
    rows: Iterable[Sequence[str]],
    widths: Sequence[int],
    *,
    font_size: float = 8.5,
) -> object:
    if sum(widths) != TOTAL_TABLE_DXA:
        raise ValueError(f"table widths must sum to {TOTAL_TABLE_DXA}: {widths}")
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    set_table_geometry(table, widths)

    header = table.rows[0]
    set_repeat_table_header(header)
    prevent_row_split(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, LIGHT_GREY)
        set_cell_text(cell, text, bold=True, color=NAVY, size=font_size)
    keep_row_with_next(header)

    for row_values in rows:
        row = table.add_row()
        prevent_row_split(row)
        for index, text in enumerate(row_values):
            cell = row.cells[index]
            set_cell_width(cell, widths[index])
            set_cell_text(cell, str(text), size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_callout(
    doc: Document,
    label: str,
    text: str,
    *,
    fill: str = PALE_BLUE,
    label_color: str = NAVY,
) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_geometry(table, [TOTAL_TABLE_DXA], indent_dxa=180)
    cell = table.cell(0, 0)
    set_cell_width(cell, TOTAL_TABLE_DXA)
    set_cell_shading(cell, fill)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    label_run = p.add_run(f"{label}  ")
    label_run.bold = True
    label_run.font.name = "Calibri"
    label_run.font.size = Pt(10)
    label_run.font.color.rgb = RGBColor.from_string(label_color)
    text_run = p.add_run(text)
    text_run.font.name = "Calibri"
    text_run.font.size = Pt(10)
    set_cell_margins(cell, 140, 180, 140, 180)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(doc: Document, lines: Sequence[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_geometry(table, [TOTAL_TABLE_DXA], indent_dxa=180)
    cell = table.cell(0, 0)
    set_cell_width(cell, TOTAL_TABLE_DXA)
    set_cell_shading(cell, "F7F9FB")
    cell.text = ""
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string("263238")
    set_cell_margins(cell, 140, 180, 140, 180)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_bullets(doc: Document, items: Sequence[str], level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for item in items:
        paragraph = doc.add_paragraph(style=style)
        paragraph.add_run(item)


def add_numbered(doc: Document, items: Sequence[str]) -> None:
    numbering = doc.part.numbering_part.element
    existing_ids = [
        int(node.get(qn("w:numId")))
        for node in numbering
        if node.tag == qn("w:num")
    ]
    num_id = max(existing_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), "7")
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)

    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.extend([ilvl, num_id_node])
        p_pr.append(num_pr)
        paragraph.add_run(item)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.line_spacing = 1.10

    heading_specs = {
        "Title": (28, NAVY, 0, 10),
        "Subtitle": (14, DARK_GREY, 0, 14),
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, NAVY, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Bullet 2", "List Number"]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(3)

    if "Small Label" not in styles:
        label = styles.add_style("Small Label", WD_STYLE_TYPE.PARAGRAPH)
        label.font.name = "Calibri"
        label.font.size = Pt(8)
        label.font.bold = True
        label.font.color.rgb = RGBColor.from_string(BLUE)
        label.paragraph_format.space_after = Pt(6)
        label.paragraph_format.keep_with_next = True


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)
    section.different_first_page_header_footer = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = hp.add_run(
        f"NC25OL  •  v{SPEC_VERSION_MAJOR_MINOR}"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(DARK_GREY)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = fp.add_run("REFERENCE PACKAGE  •  NOT PRODUCTION CONFIGURATION  •  ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(DARK_GREY)
    add_page_number(fp)

    first_footer = section.first_page_footer
    first_fp = first_footer.paragraphs[0]
    first_fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = first_fp.add_run("REFERENCE PACKAGE  •  NOT PRODUCTION CONFIGURATION")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(DARK_GREY)


def add_title_page(doc: Document) -> None:
    p = doc.add_paragraph(style="Small Label")
    p.add_run("INTEGRATION SPECIFICATION")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    border.append(bottom)
    p_pr.append(border)

    doc.add_paragraph("NC25OL", style="Title")
    doc.add_paragraph("Navigational Cybernetics 2.5 Open Ledger", style="Subtitle")
    doc.add_paragraph(
        "Domain-neutral connection contract with three reference profiles",
        style="Subtitle",
    )

    add_table(
        doc,
        ["Document", "Value"],
        [
            ("Version", SPEC_VERSION),
            ("Date", SPEC_DATE_TEXT),
            ("Author", "Maksim Barziankou (MxBv)"),
            ("Affiliation", "The Urgrund Laboratheory"),
            ("Licence", "CC BY 4.0 (specification) • MIT (code and schemas)"),
            ("Status", "Reference and conformance package"),
            ("Core anchor", "NC2.5 v3.0 • DOI 10.17605/OSF.IO/NHTC5"),
            (
                "Core SHA-256",
                "20F1EA17E0B986627CAD9AFA7C02E1DC5088C43A779435E5A38F11E43A06C9C3",
            ),
        ],
        [2400, 6960],
        font_size=9,
    )

    add_callout(
        doc,
        "DECISION",
        "Use one universal protocol for connection mechanics, but require an explicit domain profile and instance declaration before any system can become executable.",
        fill=LIGHT_BLUE,
    )

    doc.add_heading("What a receiving institution can determine", level=2)
    add_bullets(
        doc,
        [
            "which exact objects must be declared and signed before activation;",
            "which role may activate, submit, evaluate, consume, observe, or audit;",
            "which runtime sequence produces a single-use permit;",
            "which failures remain non-executable;",
            "which evidence stays local and which hashes may enter a registry;",
            "which demonstration components must be replaced before production.",
        ],
    )

    add_callout(
        doc,
        "CLAIM BOUNDARY",
        "This package is an engineering reference for domains expressible by its declared profile and runtime contract. It does not automatically classify an arbitrary deployment as NC2.5-conformant and does not replace domain policy, legal approval, or production assurance.",
        fill=AMBER,
        label_color="7F6000",
    )
    doc.add_page_break()


def add_document_map(doc: Document) -> None:
    doc.add_heading("Document map", level=1)
    rows = [
        ("1", "Executive decision", "What is universal and what remains domain-specific"),
        ("2", "Relationship to NC2.5", "Exact source anchor and claim boundary"),
        ("3", "Architecture", "Objects, layers, trust boundaries, and data flow"),
        ("4", "Roles and visibility", "Governance, operator, executor, architect, observer"),
        ("5", "Contract objects", "Profile, declaration, authority, intent, evidence, permit, receipt"),
        ("6", "Runtime state machine", "Preflight, binary gate, reservation, commit"),
        ("7", "Integrity and revision", "Hashes, idempotency, revocation, append-only state"),
        ("8", "Registry projection", "Local full evidence and shared hash-only record"),
        ("9", "Reference profiles", "Banking, document-release, and OTCS registry mappings"),
        ("10", "API contract", "Role-separated routes and security requirements"),
        ("11", "Failure matrix", "Deterministic non-executable outcomes"),
        ("12", "Production substitutions", "Controls required beyond the reference engine"),
        ("13", "Acceptance", "Local checks and go-live evidence"),
        ("14", "New-domain authoring", "How to add another connector without transfer-by-similarity"),
        ("A–C", "Appendices", "Core crosswalk, package files, synthetic walkthrough"),
    ]
    add_table(
        doc,
        ["§", "Section", "Reader outcome"],
        rows,
        [700, 2860, 5800],
        font_size=8.5,
    )


def add_executive_decision(doc: Document) -> None:
    doc.add_heading("1. Executive decision", level=1)
    doc.add_paragraph(
        "A universal connection ledger is feasible when universality is confined to the protocol boundary and the domain is expressible by the declared contract. The same lifecycle, hashes, roles, state machine, refusal posture, reservation semantics, and receipt chain can serve such domains. The domain selector cannot be universalized by assumption: every domain must declare how its actions map to effects, resources, controls, evidence, and a mapping witness."
    )
    add_table(
        doc,
        ["Universal protocol layer", "Domain-specific obligation"],
        [
            ("Canonical JSON and SHA-256 bindings", "Action alphabet and effect classes"),
            ("Profile and declaration versioning", "Instance scope and downstream targets"),
            ("Authority, revocation, and role scopes", "Prerequisites, hard blocks, and review flags"),
            ("Binary gate and non-executable dispositions", "Resource limits and trusted evidence"),
            ("Reservation and exactly-once permit consumption", "Mapping witness and approval evidence"),
            ("Append-only local ledger and receipt roots", "Production identity, atomicity, and retention"),
        ],
        [4380, 4980],
    )
    add_callout(
        doc,
        "RULE",
        "Connection similarity never creates admissibility. A new profile must carry its own mapping witness even when it reuses the same API or adapter.",
    )

    doc.add_heading("1.1 What this package is", level=2)
    add_bullets(
        doc,
        [
            "a domain-neutral connection contract;",
            "a strict schema envelope for profile, declaration, and runtime messages;",
            "a reference state machine implementing fail-closed behavior;",
            "banking, document-release, and OTCS registry profiles sharing one admission and receipt contract;",
            "a local verification set for positive, contradictory, and deliberately broken cases.",
        ],
    )
    doc.add_heading("1.2 What this package is not", level=2)
    add_bullets(
        doc,
        [
            "a payment rail, core-banking system, policy engine, or legal approval;",
            "an automatic proof that a concrete deployment belongs to an NC2.5 class;",
            "a central store for raw intent, evidence, personal data, or gate internals;",
            "a production identity, cryptographic key, time, or transactional system.",
        ],
    )


def add_nc25_relationship(doc: Document, crosswalk: dict) -> None:
    doc.add_heading("2. Relationship to NC2.5", level=1)
    doc.add_paragraph(
        "The package is anchored to one read-only NC2.5 v3.0 source file. The source hash is repeated in the profile, declaration, source manifest, and crosswalk. The package translates declared interface and separation constraints into an engineering contract; it does not modify the core or turn conditional structural results into deployment guarantees."
    )
    add_table(
        doc,
        ["Anchor", "Value"],
        [
            ("File", "NC25_v3_0_patched_v4.tex"),
            ("Version", "3.0"),
            ("DOI", "10.17605/OSF.IO/NHTC5"),
            (
                "SHA-256",
                "20F1EA17E0B986627CAD9AFA7C02E1DC5088C43A779435E5A38F11E43A06C9C3",
            ),
            ("Bytes / lines", "1,414,184 bytes / 14,228 lines"),
            ("Package posture", "Read-only input; no source modification"),
        ],
        [2300, 7060],
    )

    doc.add_heading("2.1 Load-bearing translations", level=2)
    selected = [
        entry
        for entry in crosswalk["entries"]
        if entry["id"]
        in {
            "CORE-DECLARATION",
            "CORE-NON-ACTIONABILITY",
            "A29-A33-GATE",
            "A64-A66-SEPARATION",
            "T67-HALT",
            "T98-FRAMEWORK-CORE",
            "T107-T108-TRANSFER",
            "L47-NON-WEAPONIZABILITY",
        }
    ]
    add_table(
        doc,
        ["Core anchor", "Lines", "Engineering consequence"],
        [
            (
                entry["id"],
                entry["source_lines"],
                entry["universal_use"],
            )
            for entry in selected
        ],
        [2100, 1700, 5560],
        font_size=8,
    )
    add_callout(
        doc,
        "NON-ACTIONABILITY",
        "The operator receives only a normalized permit or non-executable result. Gate geometry, margins, thresholds, sub-verdicts, reason vectors, and retained statistics stay outside operation selection.",
    )

    doc.add_heading("2.2 Accounting registers and their limits", level=2)
    add_bullets(
        doc,
        [
            "The hash-chained event ledger is an immutable append-only record of the declared event types: genealogy, not load. Permit-state changes that append no event (reservation expiry, window-limit invalidation) are visible in permit state only.",
            "The structural budget is a monotone committed-action counter read by the gate. It debits committed actions only; standing pressure, refusals, and failed commits are not accounted, so it is narrower than the core's structural-pressure accounting.",
            "Resource limits are sliding windows; committed amounts age out, so they are rate boundaries, not burden.",
            "A new declaration version starts a new accounting instance. No burden crosses versions, and a cross-version continuation claim requires an external ledger.",
            "The declared capacity carries a provenance coding (fully stipulated, empirically calibrated, derived from a load model, or mixed). The coding is sealed by the declaration hash and never read by the gate.",
        ],
    )


def add_architecture(doc: Document) -> None:
    doc.add_heading("3. Architecture", level=1)
    doc.add_heading("3.1 Seven contract layers", level=2)
    add_table(
        doc,
        ["Layer", "Object", "Purpose", "Write authority"],
        [
            ("1", "Domain profile", "Types actions, effects, resources, controls, evidence, and witness", "Governance only; new version"),
            ("2", "Connection declaration", "Binds one connector and exact instance scope to one profile hash", "Governance only; new version"),
            ("3", "Authority grant", "Limits one workload subject by action, target, and time", "Authority issuer"),
            ("4", "Intent + evidence", "Binds one proposed action to hashes, state, derived controls, and resolved evidence", "Operator proposes; production adapter resolves evidence and controls"),
            ("5", "Independent evaluator", "Runs preflight and the binary structural/domain-resource-admissibility conjunction", "Read-only over profile and declaration"),
            ("6", "Reservation + permit", "Temporarily reserves capacity and authorizes one exact commit", "Evaluator issues; executor consumes once"),
            ("7", "Local ledger + registry", "Keeps full local evidence and a minimal shared hash projection", "Append-only local writer"),
        ],
        [650, 1750, 4550, 2410],
        font_size=8,
    )

    doc.add_heading("3.2 Runtime data flow", level=2)
    add_code_block(
        doc,
        [
            "PROFILE + DECLARATION  ──sealed hashes──►  GOVERNANCE ACTIVATION",
            "AUTHORITY GRANT       ──scoped identity─►  OPERATOR INTENT + EVIDENCE",
            "INTENT                ──read-only───────►  PREFLIGHT",
            "PREFLIGHT PASS        ────────────────►  structural_gate AND effect_gate  # effect_gate checks domain-resource admissibility; it does not execute or simulate the external effect",
            "1 AND 1               ────────────────►  RESERVATION + SINGLE-USE PERMIT",
            "EXECUTOR              ──same binding──►  COMMIT / FAILED / ABORTED RECEIPT",
            "ALL EVENTS            ────────────────►  SIGNED LOCAL CHAIN",
            "HASHES + STATUS ONLY  ────────────────►  OPTIONAL SHARED REGISTRY",
        ],
    )

    doc.add_heading("3.3 Trust boundaries", level=2)
    add_bullets(
        doc,
        [
            "Connector-supplied evidence status and control booleans are untrusted claims; the production evaluator resolves and derives them from authoritative systems.",
            "Resolved evidence and derived controls are bound to the exact request and current-state anchor before gate evaluation.",
            "The connector cannot activate, revise, or write the profile or declaration.",
            "The evaluator does not expose its internal gate trace to the operator.",
            "The executor cannot consume a refusal, hold, manual-review result, or expired permit.",
            "The architect store is post-event and cannot feed a score back into action selection.",
            "The observer sees normalized event messages only and has no mutating scope.",
            "A registry projection cannot expand beyond hashes, status, revocation, and receipt roots.",
        ],
    )


def add_roles(doc: Document) -> None:
    doc.add_heading("4. Roles and visibility", level=1)
    add_table(
        doc,
        ["Role", "OAuth scope", "May do", "Must not do"],
        [
            ("Governance", "nc25.governance", "Activate declarations; issue and revoke grants; rotate and revoke permit signing keys", "Submit operator intent; consume permit; override a live refusal"),
            ("Operator", "nc25.operator", "Submit bound intent and evidence; receive minimal result", "Read gate internals; write profile, declaration, or architect trace"),
            ("Executor", "nc25.executor", "Consume one valid permit against the same binding and state", "Execute without permit; reuse permit; change action or target"),
            ("Architect", "nc25.architect", "Read post-event decisions, permits, chain head, and registry projection", "Feed trace data into action selection or alter the live gate"),
            ("Observer", "No authority principal", "Read normalized channel, message code, event time, and subject reference", "Activate, evaluate, consume, or access architect evidence"),
        ],
        [1250, 1750, 3150, 3210],
        font_size=8,
    )
    add_callout(
        doc,
        "SEPARATION",
        "mTLS authenticates the connection; scoped workload identity authorizes the role. The four authority scopes are non-interchangeable.",
    )


def add_contract_objects(doc: Document) -> None:
    doc.add_heading("5. Contract objects", level=1)
    doc.add_heading("5.1 Domain profile", level=2)
    add_table(
        doc,
        ["Field group", "Required content", "Why it is load-bearing"],
        [
            ("Core anchor", "File, version, DOI, SHA-256", "Prevents silent rebinding to another source"),
            ("Action alphabet", "Action ID, description, effect class, structural cost, resources", "Closes the candidate type"),
            ("Effect classes", "Effect ID and external-mutation flag", "Separates action from effect"),
            ("Resource catalog", "Unit, claim rule, reservation requirement", "Makes domain capacity explicit"),
            ("Controls", "Prerequisites, hard blocks, review flags, safe fallback", "Separates gate logic from workflow routing"),
            ("Evidence requirements", "Action coverage and evidence types", "Prevents unsupported mapping"),
            ("Mapping witness", "Method, evidence, full action/effect coverage, hash", "Blocks transfer-by-similarity"),
            ("Roles and failure posture", "Scopes, forbidden operator fields, non-executable outcomes", "Preserves interface separation"),
        ],
        [1900, 4250, 3210],
        font_size=8,
    )

    doc.add_heading("5.2 Connection declaration", level=2)
    add_bullets(
        doc,
        [
            "binds one profile identifier, version, and canonical hash;",
            "identifies one connector, identity issuer, and permitted environments;",
            "narrows actions, effects, targets, explicit action-target bindings, and scope values;",
            "declares a structural selector, capacity, and capacity provenance coding independently from domain limits;",
            "declares resource limits per intent and per time window;",
            "names owners, interfaces, off-limits systems, and change process;",
            "binds the mapping witness and immutable approval evidence.",
        ],
    )

    doc.add_heading("5.3 Runtime objects", level=2)
    add_table(
        doc,
        ["Object", "Binding", "Executable?"],
        [
            ("Authority grant", "Profile + declaration hashes, subject, connector, actions, targets, time", "No"),
            ("Intent", "Grant, profile + declaration hashes, action, target, scope, state, resource claims", "No"),
            ("Evidence bundle", "Intent ID, typed references, content hashes, evaluator-resolved status", "No"),
            ("Operator result", "Intent and request hashes; normalized disposition", "Only ALLOW carries a permit hash"),
            ("Execution permit", "Exact request, state, grant, action, effect, target, reservation, expiry", "Yes, once"),
            ("Execution receipt", "Permit, outcome, debits, local event hash, receipt hash", "Post-event evidence"),
        ],
        [1700, 5770, 1890],
        font_size=8,
    )

    doc.add_heading("5.4 Canonical hash rule", level=2)
    add_code_block(
        doc,
        [
            "bytes = UTF8(JSON(value, sort_keys=true, separators=(',', ':'), ensure_ascii=false))",
            "object_hash = UPPERCASE_HEX(SHA256(bytes))",
            "witness_hash = SHA256(canonical witness object with witness_hash omitted)",
        ],
    )
    doc.add_paragraph(
        "A changed canonical byte sequence is a changed object. A material change must therefore produce a new version, new hash, and new approval record."
    )


def add_state_machine(doc: Document) -> None:
    doc.add_heading("6. Runtime state machine", level=1)
    doc.add_heading("6.1 Declaration lifecycle", level=2)
    add_table(
        doc,
        ["State", "Entry condition", "Permitted transition"],
        [
            ("DRAFT", "Profile and declaration validate but are not activated", "ACTIVE after governance approval of exact hashes"),
            ("ACTIVE", "Reference status is ACTIVE, current time is valid, and the chain is intact", "No additional reference-engine lifecycle command"),
            ("OUTSIDE VALIDITY (guard)", "Current time is before effective_from or at/after expires_at", "Execution is blocked; use a new approved version or a production governance transition"),
        ],
        [1200, 4710, 3450],
        font_size=8,
    )
    doc.add_paragraph(
        "The reference engine implements DRAFT to ACTIVE plus validity blocking. SUSPENDED, REVOKED, and EXPIRED are production registry projection states whose transitions must be defined and persisted by a production governance adapter."
    )

    doc.add_heading("6.2 Evaluation order", level=2)
    add_numbered(
        doc,
        [
            "Verify profile seal, declaration seal, and local ledger chain.",
            "Verify active validity, strict message shape, and forbidden operator inputs.",
            "Resolve idempotency from the canonical connector-submitted request and reject a changed request under the same key.",
            "Verify declaration binding, connector, grant, subject, time, explicit action-target binding, scope, and revocation.",
            "Resolve authoritative evidence, verify provenance and hashes, and independently derive control results.",
            "Compute the current structural-capacity input after active reservations without emitting a verdict.",
            "Compute the domain-resource and window input without emitting a verdict.",
            "Engage the binary conjunction: the structural branch decides capacity and the effect branch decides resource admissibility.",
            "Issue a short-lived reservation and permit only for one-and-one.",
        ],
    )
    doc.add_paragraph(
        "The SDK resolves exact submitted-request idempotency first, verifies binding, authority, action-target, and scope authorization, and only then calls the injected evidence and control resolvers. Connector claims are ignored on each wired surface, resolver output is validated and bound into the result and permit request hash, and resolver failure is closed. Without the resolvers, evaluate_intent gates on the declared intent flags and remains the explicit internal post-resolution boundary — a documented reference boundary that production closes by wiring both resolvers."
    )

    doc.add_heading("6.3 Gate semantics", level=2)
    add_code_block(
        doc,
        [
            "structural_gate ∈ {0,1}",
            "effect_gate     ∈ {0,1}",
            "gate_bit = structural_gate AND effect_gate  # effect_gate checks domain-resource admissibility; it does not execute or simulate the external effect",
            "",
            "gate_bit = 1  ⇒  ALLOW + reservation + single-use permit",
            "gate_bit = 0  ⇒  REFUSAL; no permit; no execution path",
        ],
    )
    add_table(
        doc,
        ["Disposition", "Channel", "Permit", "Executor behavior"],
        [
            ("ALLOW", "OUTPUT", "Required", "May consume once before expiry against same state"),
            ("REFUSAL", "REFUSAL", "None", "Must not execute"),
            ("MANUAL_REVIEW", "REFUSAL", "None", "Route outside automated execution"),
        ],
        [1650, 1450, 1350, 4910],
        font_size=8,
    )
    doc.add_paragraph(
        "The disposition alphabet is exactly these three values, and every one is producible by the reference engine. Integrity failures and availability faults halt as transport-level errors (HTTP 503 or 400); they are never surfaced as operator results, so there is no declared disposition the executable reference cannot emit."
    )

    doc.add_heading("6.4 Minimal operator projection", level=2)
    add_code_block(
        doc,
        [
            "{",
            '  "intent_id": "...",',
            '  "request_hash": "...",',
            '  "disposition": "ALLOW | REFUSAL | MANUAL_REVIEW",',
            '  "channel": "OUTPUT | REFUSAL",',
            '  "message_code": "EXECUTION_AUTHORIZED | EXECUTION_DENIED | HUMAN_REVIEW_REQUIRED",',
            '  "permit_hash": "... | null",',
            '  "expires_at": "... | null"',
            "}",
        ],
    )
    doc.add_paragraph(
        "The operator message code is fixed by disposition: ALLOW maps to EXECUTION_AUTHORIZED, REFUSAL to EXECUTION_DENIED, and MANUAL_REVIEW to HUMAN_REVIEW_REQUIRED. It never identifies the failed gate, threshold, margin, or limit. Detailed reason codes remain architect-only."
    )
    doc.add_paragraph(
        "The disposition itself is an unavoidable one-bit decision channel. A production adapter MUST rate-limit distinct probes per connector and alert on systematic boundary-seeking sequences."
    )


def add_integrity_revision(doc: Document) -> None:
    doc.add_heading("7. Integrity, reservation, and revision", level=1)
    doc.add_heading("7.1 Reservation and commit", level=2)
    add_bullets(
        doc,
        [
            "An ALLOW always reserves the structural cost, and reserves each domain resource claim whose resource declares reservation_required.",
            "The configured permit TTL is sealed into the permit, cannot exceed 900 seconds, and expires_at is capped by the authority-grant and declaration expiries.",
            "Trusted-time causality permits at most five seconds of forward skew for intent receipt, evidence collection, revocation, and execution commit; execution may not precede permit issuance by more than five seconds.",
            "A resource declared with reservation_required false holds nothing while a permit is outstanding; only committed amounts count against its window limit.",
            "Active reservations count against later evaluations and prevent over-issue.",
            "Immediately before COMMITTED, the engine recomputes committed usage for every claimed resource against the trusted current window. When unreserved permits race, the first admissible commit wins and a later permit is invalidated with RESOURCE_WINDOW_LIMIT.",
            "The committed window is closed at its lower bound: a debit timestamped exactly at now minus window_seconds still counts and ages out only after that boundary.",
            "COMMITTED records the declared debits and consumes the permit.",
            "FAILED or ABORTED consumes the permit but releases the reservation without debit.",
            "The executor must present the same connector, profile hash, declaration hash, and state anchor.",
            "An expired, changed, revoked, or tampered permit blocks a new commit. After a successful commit, an exact canonical retry returns the stored receipt without another event or debit, including after permit, grant, or declaration expiry or revocation; a changed retry conflicts. This read-only replay cannot authorize new work.",
        ],
    )

    doc.add_heading("7.2 Idempotency", level=2)
    doc.add_paragraph(
        "The idempotency key binds a canonical hash of the connector-submitted intent and evidence bundle before external resolution. An exact evaluation replay therefore returns an active permit or the original non-executable result without calling the resolver again; the separately computed result request hash binds the authoritative resolved bundle used by the gate. Replay of a consumed, expired, revoked, or otherwise invalidated permit fails closed and never reissues ALLOW. A different submitted request under the same key is a conflict and produces no second decision or permit. Executor replay is separate: an exact canonical retry after a successful commit returns the stored receipt without a second event or debit, including after permit, grant, or declaration expiry or revocation, while a changed retry conflicts. This read-only replay cannot authorize new work. In the HTTP adapter, the Idempotency-Key header must exactly equal intent.idempotency_key; a mismatch is rejected before evaluation."
    )

    doc.add_heading("7.3 Revocation", level=2)
    doc.add_paragraph(
        "A grant revocation recorded before commit invalidates every outstanding permit and releases its reservation. A later grant cannot revive an old permit. Revocation evidence is appended to the local chain. A caller-supplied revoked_at may be at most five seconds ahead of the trusted processing clock."
    )

    doc.add_heading("7.4 Append-only local evidence", level=2)
    add_bullets(
        doc,
        [
            "Every event carries an index, UTC time, event type, role, subject reference, payload, previous-event hash, event hash, and signature.",
            "A profile, declaration, permit, signature, or chain mismatch halts execution.",
            "Raw secrets and reusable credentials are forbidden even in the local reference ledger.",
            "Production storage must provide transactional append-only or WORM behavior and independent restore evidence.",
        ],
    )

    doc.add_heading("7.5 Material revision", level=2)
    add_table(
        doc,
        ["Changed item", "Required response"],
        [
            ("Action, effect, or resource mapping", "New profile version, witness hash, profile hash, and approval"),
            ("Control flag or evidence requirement", "New profile version and affected declarations"),
            ("Connector, target, scope, limit, owner, or interface", "New declaration version and approval"),
            ("Grant scope or validity", "New grant; old permit remains bound to old grant"),
            ("Runtime observation only", "Append evidence; do not rewrite the live gate"),
        ],
        [3300, 6060],
    )


def add_registry(doc: Document) -> None:
    doc.add_heading("8. Registry projection", level=1)
    doc.add_paragraph(
        "The universal protocol does not require one central evidence ledger. Each connection retains its full evidence locally. An optional shared registry provides discoverability and revocation without becoming a source of operator signals or sensitive evidence."
    )
    add_table(
        doc,
        ["Local evidence store", "Shared registry"],
        [
            ("Profile and declaration approved bytes", "Profile and declaration identifiers, versions, and hashes"),
            ("Authority, intent, evidence references, permit, decision trace", "Connector identifier and core anchor"),
            ("Reservations, execution requests, receipts, chain events", "Status, activation time, revocation flag"),
            ("Architect-only gate sub-verdicts and reason codes", "Local ledger head hash and latest receipt root"),
            ("Access-controlled operational evidence", "No raw intent, evidence, personal data, credentials, or gate internals"),
        ],
        [4680, 4680],
    )
    add_callout(
        doc,
        "MINIMIZATION",
        "A registry expansion that introduces raw intent, raw evidence, gate geometry, sub-verdicts, reason vectors, personal data, reusable credentials, or control-plane configuration is a contract failure.",
    )


    doc.add_paragraph(
        'In the reference projection, status is the declaration lifecycle state. revoked is a connector-level alert that one or more grant revocations exist; it is not a connector shutdown bit or an authorization source. Exact grant validity remains local and is rechecked at evaluation and commit.'
    )

def add_banking_profile(doc: Document, questionnaire_crosswalk: dict) -> None:
    doc.add_heading("9. Reference profiles", level=1)
    doc.add_heading("9.1 Banking questionnaire mapping", level=2)
    doc.add_paragraph(
        "The banking profile is the first concrete use of the universal contract. It expresses a blank bank onboarding questionnaire as typed profile and declaration fields. The synthetic values demonstrate behavior only; a bank must replace them with exact enterprise identifiers, thresholds, owners, and immutable evidence."
    )
    add_table(
        doc,
        ["Questionnaire section", "Universal destination", "Execution consequence"],
        [
            (
                item["questionnaire_section"],
                "; ".join(item["universal_paths"]),
                item["connection_effect"],
            )
            for item in questionnaire_crosswalk["sections"]
        ],
        [2500, 3550, 3310],
        font_size=7.6,
    )

    doc.add_heading("9.2 Reference banking actions", level=2)
    add_table(
        doc,
        ["Action", "Effect", "Structural cost", "Domain resource"],
        [
            ("BANK_PAYMENT_POST", "BANK_LEDGER_MUTATION", "100", "money_minor_units"),
            ("BANK_PAYMENT_STATUS_READ", "BANK_READ_ONLY_QUERY", "10", "query_units"),
            ("NO_EFFECT_FALLBACK", "NO_EFFECT", "0", "None"),
        ],
        [2600, 2700, 1600, 2460],
    )
    add_callout(
        doc,
        "BANK LIMITS",
        "Per-payment and aggregate currency limits are banking resource controls. They are not the NC2.5 structural selector budget.",
        fill=AMBER,
        label_color="7F6000",
    )

    doc.add_heading("9.3 Synthetic banking allowed path", level=2)
    add_numbered(
        doc,
        [
            "Activate the banking profile and synthetic connection declaration.",
            "Issue the synthetic workload grant for the declared actions and targets.",
            "Submit BANK_PAYMENT_POST with 12,500 GBP minor units, exact scope, current-state hash, and three evaluator-resolved valid evidence items.",
            "Verify structural capacity and the domain resource window.",
            "Receive ALLOW with a permit hash and expiry; no gate internals are returned.",
            "Consume the permit once with the same binding and state anchor.",
            "Record a structural debit of 100, a domain debit of 12,500, and a signed receipt.",
        ],
    )

    doc.add_heading("9.4 Document-release reference profile", level=2)
    doc.add_paragraph(
        "The second complete profile maps an approved document release to the same protocol mechanics without importing banking actions or limits. It demonstrates evaluate-to-commit reuse for another declared domain; it is not a proof of arbitrary domain admission."
    )
    add_table(
        doc,
        ["Action", "Effect", "Structural cost", "Domain resource"],
        [
            ("DOCUMENT_RELEASE", "DOCUMENT_STATE_MUTATION", "40", "document_units"),
            ("NO_EFFECT_FALLBACK", "NO_EFFECT", "0", "None"),
        ],
        [2600, 3000, 1600, 2160],
    )


    doc.add_heading("9.5 OTCS registry bridge", level=2)
    doc.add_paragraph(
        "The OTCS bridge is built inside this ledger. The active NC2.5 profile and declaration define the admissibility atmosphere, and the NC2.5 evaluator alone issues ALLOW, REFUSAL, or MANUAL_REVIEW. OTCS does not supply an admissibility verdict: it supplies typed rights, consent, operating-grant, mark-permission, and expected-head facts, then performs a compare-and-swap append only under a live NC2.5 permit."
    )
    add_table(
        doc,
        ["Boundary", "Owner", "Invariant"],
        [
            ("Admissibility", "NC2.5 ledger", "Only the active profile/declaration and local evaluator can issue an executable permit"),
            ("Domain facts", "OTCS", "RFC 8785 digests and the expected project head are treated as typed, opaque evidence"),
            ("External mutation", "OTCS", "One project-scoped idempotent compare-and-swap append under the supplied permit"),
            ("Execution receipt", "NC2.5 ledger", "The exact OTCS receipt digest is sealed into the EXECUTION_RECORDED event"),
            ("Crash recovery", "Local SQLite outbox", "A dead PREPARED permit stops before OTCS; CALLING or AMBIGUOUS replays only with a live permit, otherwise RECONCILIATION_REQUIRED; committed effects are never re-authorized"),
        ],
        [2100, 1900, 5360],
        font_size=8,
    )
    add_numbered(
        doc,
        [
            "Read the current OTCS project head and typed legal receipts.",
            "Build an NC2.5 intent and evidence bundle bound to those hashes.",
            "Bind the complete append operation into the evidence submitted for local evaluation.",
            "Evaluate locally; only on ALLOW persist the permitted request with its original intent/evidence snapshot. Otherwise stop without contacting OTCS.",
            "Before execution or recovery, verify that snapshot and operation against the durable engine's authenticated submitted decision and recheck their agreement under the bridge's intent and evidence rules. Recheck every transition snapshot; exclude local admission data from the OTCS request.",
            "Immediately before a first OTCS call, revalidate the row's profile/declaration binding and permit; a dead PREPARED permit becomes NOT_EXECUTED without an external call.",
            "Ask OTCS to compare-and-swap the expected head. The production port verifies the receipt signature and inclusion proof before returning it.",
            "After an ambiguous CALLING state, replay the exact request only with a live permit. For CALLING or AMBIGUOUS with a dead permit, stop in RECONCILIATION_REQUIRED without another OTCS call.",
            "Bind the OTCS receipt digest into the local execution event; route a late local-finalization failure to reconciliation instead of inventing a new authorization.",
        ],
    )
    add_callout(
        doc,
        "AUTHORITY BOUNDARY",
        "Legacy outbox rows without an admission snapshot are refused with OTCS_OUTBOX_ADMISSION and require operator reconciliation. The binding trusts the durable engine and its authentication keys; it is not a MAC over all outbox state or a downstream receipt seal. A stale OTCS head is an execution failure against changed state, not a new OTCS admission decision. A retry must obtain a fresh state anchor and pass NC2.5 evaluation again.",
        fill=AMBER,
        label_color="7F6000",
    )


def add_api(doc: Document) -> None:
    doc.add_heading("10. API contract", level=1)
    add_table(
        doc,
        ["Method and route", "Scope", "Purpose"],
        [
            ("POST /v1/declarations:activate", "nc25.governance", "Activate exact profile and declaration hashes"),
            ("POST /v1/authority-grants", "nc25.governance", "Issue a scoped workload grant"),
            ("POST /v1/authority-revocations", "nc25.governance", "Revoke a grant and outstanding permits"),
            ("POST /v1/operator/intents:evaluate", "nc25.operator", "Evaluate intent and return minimal result"),
            ("POST /v1/executor/permits/{permit_hash}:consume", "nc25.executor", "Consume a live permit exactly once"),
            ("GET /v1/architect/permits/{permit_hash}", "nc25.architect", "Read post-event permit trace"),
            ("GET /v1/architect/decisions/{request_hash}", "nc25.architect", "Read post-event decision trace"),
            ("GET /v1/architect/ledger/head", "nc25.architect", "Verify append-only chain head"),
            ("GET /v1/registry/connectors/{connector_id}", "nc25.architect", "Read the exact connector hash-only projection or return 404"),
        ],
        [3900, 1900, 3560],
        font_size=8,
    )
    add_bullets(
        doc,
        [
            "Every route requires mTLS and a workload identity carrying the exact route scope.",
            "Every public SDK operation requires an explicit actor_scope; omission fails closed, no method supplies its required role as a default, and production derives the scope from verified workload identity.",
            "The runtime intent and evidence schemas are internal evaluator inputs; a production adapter ignores connector-asserted truth values and overwrites them with authoritative evidence resolution and independently derived controls.",
            "Resolved evidence and derived controls are bound to the request and current-state anchor before evaluation.",
            "The operator route requires an Idempotency-Key header exactly equal to intent.idempotency_key; mismatch is rejected before evaluation.",
            "Wire SHA-256 inputs use exactly 64 hexadecimal characters in either case; the engine normalizes comparisons and emits canonical uppercase values.",
            "The executor consume route requires the path permit_hash to equal execution_request.permit_hash. The adapter passes the path value as route_permit_hash to commit_execution; mismatch fails with EXECUTION_PERMIT_MISMATCH before lookup, replay, consumption, or debit.",
            "The registry route passes its connector_id path selector to registry_record. A different or unknown selector fails with CONNECTOR_NOT_FOUND and maps to HTTP 404; the active connector record is never substituted.",
            "If an in-process exception interrupts reference commit, the engine restores its pre-call state before re-raising. With state_path configured, SQLite serializes writers with BEGIN IMMEDIATE, commits permit consumption, debit, receipt, replay state, reservations, and the local event chain atomically for one node, and reloads one verified snapshot under the local lock for every public read.",
            "Signer configuration is external runtime configuration. Reopening a state_path whose event chain contains Ed25519 envelopes requires a verifier-compatible event_signer with the same key_id; SQLite deliberately persists neither signer configuration nor private keys. Omitting the verifier, or supplying one that declares both an algorithm and a key_id and differs from the persisted envelope in either, fails closed with EVENT_SIGNER_REQUIRED. The EventSigner protocol requires only sign and verify, so a signer that does not declare both cannot be compared against the envelope at all; for such a signer the mismatch is not distinguished and reaches the caller as a chain verdict. A signer declaring the same key_id with different key material likewise cannot be told apart from a forged chain and so fails full-chain integrity verification.",
            "External JSON Schema references define profile, declaration, and runtime payloads.",
            "The HTTP adapter wraps the reference engine's primitive declaration and grant hashes into the response objects defined by OpenAPI.",
            "A standard-library reference WSGI adapter (sdk/python/nc25_universal_adapter.py) implements every route and enforces these wire rules end-to-end: it derives the route scope from the transport identity and never the body, binds the posted activation profile and declaration to the engine's sealed hashes, requires the Idempotency-Key header to equal intent.idempotency_key, binds the path permit hash to the body, and maps engine outcomes to the declared HTTP statuses. The engine's default paths are standard-library only; the optional Ed25519 signers require the pinned cryptography wheel and fail closed without it. Production replaces the scope header with mTLS plus scoped OAuth.",
            "A contract error is distinct from a normalized non-executable business result.",
            "Integrity failure returns no permit and requires an operational halt.",
        ],
    )


def add_failure_matrix(doc: Document) -> None:
    doc.add_heading("11. Failure matrix", level=1)
    add_table(
        doc,
        ["Condition", "Disposition / error", "Gate engaged", "Permit"],
        [
            ("Missing required evidence", "REFUSAL", "No", "None"),
            ("Structurally valid evidence with INVALID, REVOKED, or EXPIRED status", "REFUSAL", "No", "None"),
            ("Malformed evidence object, hash, reference, or unknown status", "Contract error / halt", "No decision", "None"),
            ("Ambiguous request", "MANUAL_REVIEW", "No", "None"),
            ("Manual-review flag true", "MANUAL_REVIEW", "No", "None"),
            ("Hard-block flag true", "REFUSAL", "No", "None"),
            ("Binding, authority, action-target pairing, or scope mismatch", "REFUSAL", "No", "None"),
            ("Structural capacity unavailable", "REFUSAL", "Yes; structural=0", "None"),
            ("Domain resource rule or window unavailable", "REFUSAL", "Yes; effect=0", "None"),
            ("Idempotency key reused with changed request", "Conflict", "No new evaluation", "None"),
            ("Grant revoked before commit", "Commit rejected", "Prior result invalidated", "None usable"),
            ("State anchor changed", "Commit rejected", "Prior result invalidated", "None usable"),
            ("Permit expired, tampered, or consumed", "Commit rejected", "No new evaluation", "None usable"),
            ("Profile, declaration, or ledger integrity failure", "Transport-level halt; no operator result", "No", "None"),
        ],
        [3600, 2250, 2050, 1460],
        font_size=7.8,
    )


def add_production(doc: Document) -> None:
    doc.add_heading("12. Production substitutions and security", level=1)
    add_table(
        doc,
        ["Reference element", "Production requirement", "Acceptance evidence"],
        [
            ("Reference HMAC permit signer and state HMAC", "Asymmetric KMS/HSM key, rotation, revocation, custody", "Key policy and independent verification"),
            ("Optional Ed25519 permit/event envelope", "KMS/HSM-backed detached signature with managed key policy", "Tamper, wrong-key, rotation, and independent-verification evidence"),
            ("Versioned permit signature envelope (shipped; key_id-routed verification, rotation, revocation)", "Managed key custody source behind the same envelope; key ceremonies and revocation evidence", "Wire-contract conformance and independent signature verification"),
            ("Connector-supplied evidence claims", "Inject an authoritative evidence resolver; exceptions, malformed output, and intent mismatch fail closed", "Override, malformed-output, and request/state binding tests"),
            ("Connector-supplied control booleans", "Inject an authoritative control resolver (shipped injection point); derived results replace the intent flags and bind into the request hash; failures fail closed", "Independent control-derivation and spoofed-control-rejection tests"),
            ("In-memory event list", "Transactional append-only or WORM storage", "Retention configuration and restore test"),
            ("Optional single-node SQLite idempotency", "Managed or replicated unique constraint and replay-safe result store where multi-node availability is required", "Concurrency, restart, and failover tests"),
            ("In-memory execution replay", "Store execution request hash and receipt atomically with permit consumption", "Lost-response retry and changed-retry conflict tests"),
            ("Required local OTCS SQLite outbox", "Transactional replicated outbox or workflow store preserving exact requests, receipts, and reconciliation state", "Crash-before-call, ambiguous-response, and failover tests"),
            ("Project-idempotent OTCS append port", "Authenticated expected-head compare-and-swap client with signature and inclusion-proof verification", "Duplicate suppression, stale-head, signature, proof, and replay tests"),
            ("Optional SQLite committed_resources history", "Managed windowed aggregate with pruning or compaction, replication where required, and trusted time", "Window-boundary, restart, retention, and failover tests"),
            ("In-process RLock", "Database or distributed serialization across reservation, debit, and receipt creation", "Concurrent oversubscription and multi-process tests"),
            ("Full-chain verification per operation", "Indexed or checkpointed hot-path verification plus independent full-chain audits", "Scale benchmark and historical tamper tests"),
            ("Process clock", "Trusted UTC source with drift monitoring", "Clock source, thresholds, and alert evidence"),
            ("Process reservation", "Durable reservation coordinated with downstream commit", "Failure-injection and double-spend tests"),
            ("Function-call role", "mTLS and enterprise workload identity with scoped OAuth", "Trust chain and access matrix"),
            ("Synthetic evidence URI", "Immutable access-controlled evidence reference", "Evidence registry and hash verification"),
            ("In-process commit", "Atomic or compensatable downstream transaction boundary", "Commit/rollback design and reconciliation test"),
        ],
        [2150, 4170, 3040],
        font_size=7.8,
    )

    doc.add_heading("12.1 Security invariants", level=2)
    add_bullets(
        doc,
        [
            "Connector-supplied evidence status and prerequisite, hard-block, or manual-review booleans are never authority in production: the engine ships an evidence resolver and a control resolver whose derived results replace the connector claims as gate inputs, and production MUST wire both to authoritative systems. Without the resolvers the reference engine gates on the declared intent flags — a documented reference boundary, not a production posture.",
            "The production evaluator resolves authoritative evidence and derives controls before the reference gate boundary.",
            "No raw secret, token, password, private key, or reusable credential enters an intent, event, receipt, or registry record.",
            "The connector has no write path to profile, declaration, gate, or architect store.",
            "Operator and observer surfaces never expose gate internals.",
            "Revocation and expiry are checked again at commit, not only at evaluation.",
            "A local integrity failure blocks every later operation until investigated.",
            "Logs and evidence use enterprise data-classification, retention, and access controls.",
        ],
    )

    doc.add_heading("12.2 Contract validation and bounded failure policy", level=2)
    doc.add_paragraph(
        "failure_posture is a required profile mirror of protocol-fixed outcomes: missing_required_data is REFUSAL and ambiguity is MANUAL_REVIEW. Neither the profile nor the declaration can select an alternative disposition."
    )
    add_bullets(
        doc,
        [
            "Both posture values are required, non-executable, and fixed: missing_required_data is REFUSAL and ambiguity is MANUAL_REVIEW.",
            "Integrity failures remain raised hard errors because a compromised seal or ledger cannot safely author a new disposition.",
            "This reference has no downstream caller or independent rule-conflict detector, so integrity_failure, downstream_unavailable, and rule_conflict are not decorative configuration fields.",
            "A production extension adds a reviewed runtime branch and a new schema version before adding another configurable failure condition.",
            "Every inbound profile, declaration, and runtime message has a normative JSON Schema, exercised by the shipped conformance suite. The reference engine enforces its semantic contract itself; wire-ingress schema validation is an integrator obligation (README step 5), because the standard-library reference adapter deliberately carries no third-party schema dependency.",
            "The conformance suite activates date-time format checking and carries negative probes for malformed formats and overlong strings.",
        ],
    )



def add_acceptance(doc: Document) -> None:
    doc.add_heading("13. Acceptance", level=1)
    doc.add_heading("13.1 Reference package checks", level=2)
    add_table(
        doc,
        ["Check set", "Coverage", "Expected result"],
        [
            ("Functional behavior", f"{FUNCTIONAL_TEST_COUNT} positive and negative state-machine tests, each run against all three shipped profiles", "All pass on all three profiles"),
            ("Semantic contradictions", f"{SEMANTIC_TEST_COUNT} profile, declaration, role, and input tests", "All rejected as designed"),
            ("Deliberate-break sensitivity", f"{SAFETY_MUTATION_COUNT} named safety facts", "Each mutation produces its own expected refusal or error"),
            ("Contract conformance", f"{CONFORMANCE_CHECK_COUNT} checks covering shipped artefacts, active format and length bounds, live outputs, wire references, route table, per-route security, and engine operations", "All validate and resolve; reported NOT_RUN, never as a pass, when the optional libraries are absent"),
            ("Contract regressions", f"{EXTERNAL_REGRESSION_COUNT} focused tests derived from the executable suite", "Each pins one contract fact: a defect the engine once had, or a refusal code that no other check witnessed"),
            ("OTCS bridge", f"{OTCS_BRIDGE_TEST_COUNT} tests over the permit-gated append seam, including every declared refusal code of its own vocabulary", "Each refusal answers with its own code, and the declared vocabulary equals the measured one"),
            ("Failure-surface ratchet", "Layered closed-world vocabularies, both emission-point counts with the dynamic gap enumerated, runtime-observed witnesses, and architect decisions", "No unclassified emission path, no silent runtime sink, no unacknowledged baseline delta"),
            ("Package integrity", "Sources, schemas, API routes, hashes, registry, document", "All package checks pass"),
        ],
        [2500, 4510, 2350],
    )
    doc.add_paragraph(
        "These checks exercise the named behaviours of the supplied reference package. They do not substitute for independent production review, domain approval, or infrastructure assurance."
    )

    doc.add_heading("13.2 Minimum go-live evidence", level=2)
    add_bullets(
        doc,
        [
            "approved profile bytes, profile hash, and complete mapping witness;",
            "approved declaration bytes, declaration hash, and exact enterprise identifiers;",
            "authoritative evidence resolvers and independently derived control results bound to request and state;",
            "workload identity, mTLS trust chain, and role access matrix;",
            "KMS/HSM custody, rotation, and revocation procedure;",
            "versioned signature envelope with algorithm, key identifier, and signature bytes;",
            "trusted time source and drift monitoring;",
            "multi-node idempotency, cross-system reservation, and downstream atomicity design;",
            "append-only or WORM retention configuration and restore test;",
            "positive connection test plus negative tests for every fail-closed condition;",
            "change process that creates a new version instead of altering a live gate;",
            "independent review of domain mappings, production code, and deployment controls.",
        ],
    )
    add_callout(
        doc,
        "GO-LIVE BOUNDARY",
        "Synthetic fixtures, optional single-node SQLite durability, or a passing local test set must never be promoted as production evidence.",
        fill=RED_LIGHT,
        label_color="9C0006",
    )


def add_new_domain(doc: Document) -> None:
    doc.add_heading("14. Adding a new domain", level=1)
    add_numbered(
        doc,
        [
            "Name the domain and define a closed action alphabet.",
            "Define effect classes independently from action names.",
            "Declare every resource, unit, claim rule, and reservation requirement.",
            "Map every action to exactly one effect class and its required resources.",
            "Define prerequisites, hard blocks, review flags, and a zero-effect fallback.",
            "Define evidence requirements that cover every action.",
            "Create a mapping witness covering every action and effect class.",
            "Define role visibility, forbidden operator fields, failure posture, and hash-only registry policy.",
            "Seal the profile and create a narrower instance declaration with explicit action-target bindings.",
            "Run positive, contradiction, and deliberate-break checks before integration.",
            "If the domain cannot be represented without weakening this contract, stop and create a new reviewed schema version instead of force-fitting it.",
        ],
    )
    add_table(
        doc,
        ["Invalid shortcut", "Required repair"],
        [
            ("The new API looks like the banking API", "Provide a new domain mapping witness"),
            ("The old action name is reused", "Declare its effect class and resources in the new domain"),
            ("The adapter already has approvals", "Bind immutable evidence and exact authority to the new declaration"),
            ("A manual reviewer can override refusal", "Create a new approved version or proceed outside the automated path"),
            ("A central ledger can keep all details", "Keep full evidence local and publish only the hash-only projection"),
        ],
        [4100, 5260],
    )

    doc.add_heading("14.1 Profile validation command", level=2)
    add_code_block(
        doc,
        [
            "python -B -c \"import sys; sys.path.insert(0, 'sdk/python'); from nc25_universal_ledger import load_json, validate_profile; validate_profile(load_json(sys.argv[1])); print('PROFILE_VALID')\" profiles/banking/bank-profile.json",
        ],
    )


def add_appendices(doc: Document, crosswalk: dict) -> None:
    doc.add_heading("Appendix A. Complete NC2.5 crosswalk", level=1)
    add_table(
        doc,
        ["ID", "Lines", "Universal use", "Constraint"],
        [
            (
                entry["id"],
                entry["source_lines"],
                entry["universal_use"],
                entry["constraint"],
            )
            for entry in crosswalk["entries"]
        ],
        [1750, 1250, 3000, 3360],
        font_size=7.2,
    )

    doc.add_heading("Appendix B. Package files", level=1)
    add_table(
        doc,
        ["Path", "Purpose"],
        [
            ("README.md", "Reader entry point and claim boundary"),
            ("SOURCE-MANIFEST.json", "Immutable external-input hashes"),
            ("PACKAGE-MANIFEST.json", "SHA-256 and byte-length coverage of shipped package files"),
            ("requirements-lock.txt", "Pinned schema, YAML, and DOCX tool dependencies"),
            ("LICENSE", "Per-file documentation and code licensing"),
            ("core/nc25-core-crosswalk.json", "Exact core source anchors"),
            ("core/schemas/profile.schema.json", "Domain profile contract"),
            ("core/schemas/connection-declaration.schema.json", "Instance declaration contract"),
            ("core/schemas/runtime-contracts.schema.json", "Authority, intent, evidence, result, execution, and receipt contracts"),
            ("core/schemas/registry-record.schema.json", "Hash-only shared projection"),
            ("protocol/universal-connection.openapi.yaml", "Role-separated API"),
            ("sdk/python/nc25_universal_ledger.py", "Reference state machine"),
            ("sdk/python/nc25_universal_adapter.py", "Reference WSGI HTTP adapter enforcing the wire contract"),
            ("sdk/python/nc25_otcs_bridge.py", "Permit-gated OTCS append bridge with durable recovery outbox"),
            ("profiles/PROFILE-AUTHORING-CHECKLIST.md", "Required profile authoring gates"),
            ("profiles/banking/", "Bank profile and questionnaire crosswalk"),
            ("profiles/document-release/reference-connection.json", "Second executable reference profile"),
            ("profiles/otcs/reference-connection.json", "Third executable profile plus typed OTCS operation and receipt fixture"),
            ("examples/", "Synthetic executable fixtures"),
            ("tests/test_otcs_bridge.py", "OTCS authority-boundary and crash-seam regression cases"),
            ("tests/", "Behavior, contradiction, deliberate-break, conformance, and package checks"),
            ("tools/", "Deterministic builders for this document, its review projection, the package manifest, the acceptance report, the review bundle and the contract's error-code enum, plus the failure-surface ratchet"),
            ("UNIVERSAL-CONNECTION-RUNBOOK.md", "Integration sequence"),
            ("ACCEPTANCE-REPORT.md", "Verification status and production boundary"),
        ],
        [3900, 5460],
        font_size=8,
    )

    doc.add_heading("Appendix C. Synthetic command sequence", level=1)
    add_code_block(
        doc,
        [
            "python -B tests/run_acceptance.py",
            "python -B -m pytest -p no:cacheprovider tests  # optional collection",
            "",
            "Expected reference outcome:",
            "  functional tests .......... pass",
            "  semantic contradictions ... rejected",
            "  deliberate breaks ......... each detected by its named control",
            "  package bindings .......... pass",
        ],
    )


def build() -> Path:
    crosswalk = load_json(ROOT / "core" / "nc25-core-crosswalk.json")
    questionnaire_crosswalk = load_json(
        ROOT / "profiles" / "banking" / "questionnaire-crosswalk.json"
    )

    doc = Document()
    configure_styles(doc)
    configure_document(doc)
    doc.core_properties.title = "NC25OL Specification"
    doc.core_properties.author = "Maksim Barziankou (MxBv)"
    doc.core_properties.last_modified_by = "Maksim Barziankou (MxBv)"
    doc.core_properties.subject = "Domain-neutral connection contract"
    doc.core_properties.keywords = "NC2.5, connection ledger, integration, reference profiles"
    doc.core_properties.comments = "Reference package; not production configuration"

    add_title_page(doc)
    add_document_map(doc)
    add_executive_decision(doc)
    add_nc25_relationship(doc, crosswalk)
    add_architecture(doc)
    add_roles(doc)
    add_contract_objects(doc)
    add_state_machine(doc)
    add_integrity_revision(doc)
    add_registry(doc)
    add_banking_profile(doc, questionnaire_crosswalk)
    add_api(doc)
    add_failure_matrix(doc)
    add_production(doc)
    add_acceptance(doc)
    add_new_domain(doc)
    add_appendices(doc, crosswalk)

    final = doc.add_paragraph()
    final.alignment = WD_ALIGN_PARAGRAPH.CENTER
    final.paragraph_format.space_before = Pt(18)
    run = final.add_run("END OF SPECIFICATION")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(DARK_GREY)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    # Fixed on purpose: the builder is deterministic, and identical inputs
    # must give identical bytes, so these are NOT build dates and do not move
    # with a rebuild. The date that identifies a version is the acceptance
    # report's verification date; the bytes that identify it are the manifest.
    doc.core_properties.created = datetime(2026, 7, 28, 0, 0, 0)
    doc.core_properties.modified = datetime(2026, 8, 1, 0, 0, 0)
    doc.save(OUTPUT)
    _normalize_docx_archive(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
